import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_hyperlink(paragraph, url, text, color="004080", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def create_doc():
    doc = Document()

    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Techcrush NFT Assignment")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 30, 60)

    # Subtitle / Student Details
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Student: Harry Joseph Soronnadi  |  Network: Ethereum Sepolia  |  Symbol: TCRUSH\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    # Summary Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading Accent 1' if 'Light Shading Accent 1' in [s.name for s in doc.styles] else 'Table Grid'

    items = [
        ("GitHub Repository", "https://github.com/Bookofharry/techcrush-nft", "https://github.com/Bookofharry/techcrush-nft"),
        ("Deployed Contract (Sepolia)", "0x62DEd3949ED5EF6af1784Cace3393e504B9656D9", "https://sepolia.etherscan.io/address/0x62DEd3949ED5EF6af1784Cace3393e504B9656D9"),
        ("Deployment Transaction", "0xf4b789bcd502abd2bc6793480e7cdb77bc5819d0938502485dd74e1f7eabe22d", "https://sepolia.etherscan.io/tx/0xf4b789bcd502abd2bc6793480e7cdb77bc5819d0938502485dd74e1f7eabe22d"),
        ("Mint Transaction (Token #1)", "0xac5697902f1ff96f482a2388dc1f39f4adc864ab012ce42703ea46278877eb80", "https://sepolia.etherscan.io/tx/0xac5697902f1ff96f482a2388dc1f39f4adc864ab012ce42703ea46278877eb80"),
        ("Token Recipient / Owner", "0xc5e058e6B8BeAC6EFa667e958D8cA0E3EE36a292", "https://sepolia.etherscan.io/address/0xc5e058e6B8BeAC6EFa667e958D8cA0E3EE36a292"),
    ]

    for idx, (label, text, url) in enumerate(items):
        cell_lbl = table.rows[idx].cells[0]
        cell_val = table.rows[idx].cells[1]
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(label)
        r_l.font.name = "Arial"
        r_l.font.size = Pt(10)
        r_l.font.bold = True

        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(text)
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = RGBColor(0, 102, 204)

    doc.add_paragraph()

    # Section Helper
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    screenshots_dir = os.path.join(base_dir, "evidence", "screenshots")

    sections_data = [
        (
            "1. Smart Contract Compilation (forge build)",
            "The contract compiles cleanly with Foundry's Solc 0.8.30 compiler. All 43 files, including OpenZeppelin contracts (ERC721URIStorage, Ownable), are compiled without warnings.",
            os.path.join(screenshots_dir, "01-compilation.png")
        ),
        (
            "2. Automated Contract Testing (forge test -vv)",
            "All 12 unit tests pass successfully, validating collection identity, safe minting, sequential token IDs, metadata URI persistence, access control (onlyOwner), and ERC-721 transfer permissions.",
            os.path.join(screenshots_dir, "02-testing.png")
        ),
        (
            "3. Sepolia Contract Deployment",
            "The TechcrushNFT contract was successfully broadcast and verified on Ethereum Sepolia at address 0x62DEd3949ED5EF6af1784Cace3393e504B9656D9 with receipt status 0x1 (Success).",
            os.path.join(screenshots_dir, "03-sepolia-deployment.png")
        ),
        (
            "4. Sepolia Token Minting (Techcrush #1)",
            "Token #1 was successfully minted to recipient 0xc5e058e6B8BeAC6EFa667e958D8cA0E3EE36a292. Both ownerOf(1) and tokenURI(1) were verified on Ethereum Sepolia.",
            os.path.join(screenshots_dir, "04-sepolia-minting.png")
        ),
    ]

    for title, desc, img_path in sections_data:
        h = doc.add_heading(level=2)
        r_h = h.add_run(title)
        r_h.font.name = "Arial"
        r_h.font.size = Pt(14)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(20, 50, 100)

        p_d = doc.add_paragraph()
        r_d = p_d.add_run(desc)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(10.5)
        r_d.font.color.rgb = RGBColor(60, 60, 60)

        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(18)
            doc.add_picture(img_path, width=Inches(6.8))

    output_path = os.path.join(base_dir, "evidence", "Techcrush_NFT_Assignment_Harry_Joseph_Soronnadi.docx")
    doc.save(output_path)
    print(f"Generated: {output_path}")

if __name__ == "__main__":
    create_doc()
